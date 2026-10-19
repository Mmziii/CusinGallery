"""Task 3: xlsx / CSV exports.

Asserts what the owner actually receives: a real workbook openpyxl can
re-open (RTL, bold frozen header, Persian headers, numeric money, Jalali
text dates), a CSV with a UTF-8 BOM, the 50k-row cap, and that a staff
member without the section permission gets 403.
"""
import io

from django.contrib.auth.models import Permission
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from apps.adminui.tests.helpers import PLAIN_STATIC, make_staff, make_superuser
from apps.analytics.models import Event

from .helpers import make_customer, make_event, make_order, make_product

REPORTS = "admin_reports"


def workbook(response):
    return load_workbook(io.BytesIO(response.content))


@override_settings(**PLAIN_STATIC)
class ReportExportTests(TestCase):
    def setUp(self):
        self.owner = make_superuser(username="exp-owner", phone="+989173000001")
        self.client.force_login(self.owner)
        self.product = make_product(name="قوری چینی", price=250000)
        customer = make_customer(phone="+989173000002")
        make_order(customer, self.product, total=250000)

    def _export(self, section, fmt="xlsx", **params):
        params.update({"section": section, "format": fmt, "range": "30"})
        return self.client.get(reverse("admin_reports_export"), params)

    def test_xlsx_is_a_real_rtl_workbook_with_persian_headers(self):
        response = self._export("summary")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn(".xlsx", response["Content-Disposition"])

        sheet = workbook(response).active
        self.assertTrue(sheet.sheet_view.rightToLeft)
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertEqual(sheet["A1"].value, "شاخص")
        self.assertEqual(sheet["B1"].value, "مقدار")
        self.assertTrue(sheet["A1"].font.bold)
        self.assertEqual(sheet["A1"].fill.fgColor.rgb[-6:], "E8EAED")
        # Revenue is a REAL number with the thousands format.
        revenue_cell = sheet["B2"]
        self.assertIsInstance(revenue_cell.value, int)
        self.assertEqual(revenue_cell.number_format, "#,##0")
        self.assertGreaterEqual(sheet.column_dimensions["A"].width, 10)

    def test_chart_export_has_jalali_dates_as_text(self):
        response = self._export("chart")
        sheet = workbook(response).active
        self.assertEqual(sheet["A1"].value, "تاریخ (شمسی)")
        self.assertIsInstance(sheet["A2"].value, str)
        self.assertRegex(sheet["A2"].value, r"^\d{4}/\d{2}/\d{2}$")
        self.assertIsInstance(sheet["B2"].value, int)

    def test_csv_has_a_utf8_bom_and_persian_headers(self):
        response = self._export("summary", fmt="csv")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response["Content-Type"].startswith("text/csv"))
        self.assertTrue(response.content.startswith(b"\xef\xbb\xbf"))
        text = response.content.decode("utf-8-sig")
        self.assertIn("شاخص,مقدار", text.splitlines()[0])
        self.assertIn("درآمد", text)

    def test_every_section_has_both_downloads(self):
        from apps.analytics.exports import SECTION_TITLES

        for section in SECTION_TITLES:
            with self.subTest(section=section):
                self.assertEqual(self._export(section).status_code, 200)
                self.assertEqual(self._export(section, "csv").status_code, 200)

    def test_exports_keep_the_original_product_name(self):
        """Display is Persian-digit; the download stays the source text."""
        product = make_product(name="سایز 28 کد 9915")
        make_event(Event.Type.PRODUCT_VIEW, product=product, session_key="export-digits-01")

        sheet = workbook(self._export("products_viewed")).active
        values = [cell.value for row in sheet.iter_rows() for cell in row]
        self.assertIn("سایز 28 کد 9915", values)

        csv_text = self._export("products_viewed", fmt="csv").content.decode("utf-8-sig")
        self.assertIn("سایز 28 کد 9915", csv_text)

    def test_a_staff_member_needs_the_section_permission(self):
        staff = make_staff(username="exp-staff", phone="+989173000003")
        staff.user_permissions.add(Permission.objects.get(
            content_type__app_label="analytics", codename="view_event",
        ))
        staff = type(staff).objects.get(pk=staff.pk)
        self.client.force_login(staff)
        self.assertEqual(self._export("searches_top").status_code, 200)
        self.assertEqual(self._export("summary").status_code, 403)
        self.assertEqual(self._export("products_best").status_code, 403)

    def test_unknown_section_or_format_is_404(self):
        self.assertEqual(self._export("nope").status_code, 404)
        self.assertEqual(self._export("summary", fmt="pdf").status_code, 404)


@override_settings(**PLAIN_STATIC, ANALYTICS_EXPORT_ROW_LIMIT=2)
class ExportCapTests(TestCase):
    def test_xlsx_and_csv_stop_at_the_configured_row_limit(self):
        owner = make_superuser(username="cap-owner", phone="+989173000010")
        self.client.force_login(owner)
        for index in range(3):
            make_event(
                Event.Type.SEARCH, session_key=f"cap-session-{index:02d}",
                search_query=f"کالای {index}", search_results=0,
            )
        xlsx = self.client.get(
            reverse("admin_reports_export"),
            {"section": "searches_top", "format": "xlsx", "range": "30"},
        )
        sheet = workbook(xlsx).active
        self.assertEqual(sheet.max_row, 3)  # header + 2 rows

        csv = self.client.get(
            reverse("admin_reports_export"),
            {"section": "searches_top", "format": "csv", "range": "30"},
        )
        self.assertEqual(len(csv.content.decode("utf-8-sig").strip().splitlines()), 3)
