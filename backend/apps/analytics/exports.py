"""Excel / CSV exports for the reports page and admin changelist actions.

One implementation, used by both:
* ``xlsx_bytes`` builds a real .xlsx with openpyxl: RTL sheet, Persian
  headers, bold frozen header row, auto column widths, real numbers with
  the ``#,##0`` thousands format, Jalali dates as plain text;
* ``csv_bytes`` is UTF-8 WITH BOM, so Excel shows Persian correctly.

Everything is row-capped (``settings.ANALYTICS_EXPORT_ROW_LIMIT``,
default 50,000) so one click can never hold a worker until it times out;
callers surface a Persian notice when the cap was hit.
"""
from __future__ import annotations

import csv
import io

from django.conf import settings
from django.http import HttpResponse

from apps.adminui import jalali

XLSX_CONTENT_TYPE = (
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
)
CSV_CONTENT_TYPE = "text/csv; charset=utf-8"

#: Persian headers for every report section (also used as sheet names).
SECTION_TITLES = {
    "summary": "خلاصه",
    "funnel": "قیف فروش",
    "abandoned": "سبدهای رها‌شده",
    "products_best": "پرفروش‌ترین محصولات",
    "products_viewed": "پربازدیدترین محصولات",
    "products_lowconv": "پربازدید ولی کم‌فروش",
    "products_cart": "به سبد رفته ولی خریده نشده",
    "products_wishlist": "پربازدیدترین علاقه‌مندی‌ها",
    "searches_top": "پرجستجوترین عبارت‌ها",
    "searches_zero": "جستجوهای بی‌نتیجه",
    "customers_top": "مشتریان برتر",
    "customers_places": "سفارش‌ها بر اساس استان و شهر",
    "coupons": "کوپن‌ها",
    "chart": "فروش روزانه",
}

#: Every section key, in page order.
SECTION_KEYS = tuple(SECTION_TITLES)


def row_limit() -> int:
    try:
        return max(1, int(getattr(settings, "ANALYTICS_EXPORT_ROW_LIMIT", 50000)))
    except (TypeError, ValueError):  # pragma: no cover - defensive
        return 50000


# ---------------------------------------------------------------------------
# report section tables
# ---------------------------------------------------------------------------

def _jalali(value) -> str:
    return jalali.format_jalali(value, persian_digits=False) if value else ""


def _jalali_datetime(value) -> str:
    return jalali.format_jalali_datetime(value, persian_digits=False) if value else ""


def section_table(report: dict, key: str) -> dict:
    """Return ``{title, headers, rows, numeric}`` for one report section.

    ``numeric`` is the set of column indexes written as real numbers (so
    Excel can sum them); everything else (Jalali dates, names) is text.
    """
    title = SECTION_TITLES.get(key)
    if title is None:
        raise KeyError(key)

    summary = report["summary"]
    products = report["products"]
    customers = report["customers"]

    if key == "summary":
        return {
            "title": title,
            "headers": ["شاخص", "مقدار"],
            "rows": [
                ["درآمد (تومان)", summary["revenue"]],
                ["تعداد سفارش‌ها", summary["orders"]],
                ["سفارش‌های پرداخت‌شده", summary["paid_orders"]],
                ["میانگین سبد (تومان)", summary["aov"]],
                ["تعداد نشست‌ها", summary["sessions"]],
                ["نشست‌های منجر به پرداخت", summary["paid_sessions"]],
                ["نرخ تبدیل (٪)", summary["conversion"]],
                ["مشتریان جدید", summary["new_customers"]],
                ["مشتریان بازگشتی", summary["returning_customers"]],
            ],
            "numeric": {1},
        }

    if key == "funnel":
        return {
            "title": title,
            "headers": ["مرحله", "تعداد نشست", "افت نسبت به مرحله قبل (٪)"],
            "rows": [
                [row["label"], row["count"], row["drop"] if row["drop"] is not None else ""]
                for row in report["funnel"]
            ],
            "numeric": {1, 2},
        }

    if key == "abandoned":
        return {
            "title": title,
            "headers": ["مشتری", "تلفن", "تعداد اقلام", "مبلغ سبد (تومان)", "آخرین فعالیت"],
            "rows": [
                [
                    row["name"], row["phone"], row["item_count"], row["total"],
                    _jalali_datetime(row["last_activity"]),
                ]
                for row in report["abandoned"]
            ],
            "numeric": {2, 3},
        }

    if key == "products_best":
        return {
            "title": title,
            "headers": ["محصول", "تعداد فروش", "درآمد (تومان)"],
            "rows": [[row["name"], row["qty"], row["revenue"]] for row in products["best_sellers"]],
            "numeric": {1, 2},
        }

    if key == "products_viewed":
        return {
            "title": title,
            "headers": ["محصول", "تعداد بازدید"],
            "rows": [[row["name"], row["count"]] for row in products["most_viewed"]],
            "numeric": {1},
        }

    if key == "products_lowconv":
        return {
            "title": title,
            "headers": ["محصول", "تعداد بازدید", "تعداد فروش", "نرخ تبدیل (٪)"],
            "rows": [
                [row["name"], row["views"], row["sold"], row["conversion"]]
                for row in products["low_conversion"]
            ],
            "numeric": {1, 2, 3},
        }

    if key == "products_cart":
        return {
            "title": title,
            "headers": ["محصول", "افزودن به سبد", "تعداد فروش", "فاصله"],
            "rows": [
                [row["name"], row["cart_adds"], row["sold"], row["gap"]]
                for row in products["cart_not_bought"]
            ],
            "numeric": {1, 2, 3},
        }

    if key == "products_wishlist":
        return {
            "title": title,
            "headers": ["محصول", "تعداد افزودن به علاقه‌مندی"],
            "rows": [[row["name"], row["count"]] for row in products["most_wishlisted"]],
            "numeric": {1},
        }

    if key == "searches_top":
        return {
            "title": title,
            "headers": ["عبارت جستجو", "تعداد"],
            "rows": [[row["query"], row["count"]] for row in report["searches"]["top"]],
            "numeric": {1},
        }

    if key == "searches_zero":
        return {
            "title": title,
            "headers": ["عبارت جستجو", "تعداد"],
            "rows": [[row["query"], row["count"]] for row in report["searches"]["zero"]],
            "numeric": {1},
        }

    if key == "customers_top":
        return {
            "title": title,
            "headers": ["مشتری", "تلفن", "تعداد سفارش", "مجموع خرید (تومان)"],
            "rows": [
                [row["name"], row["phone"], row["orders"], row["spend"]]
                for row in customers["top"]
            ],
            "numeric": {2, 3},
        }

    if key == "customers_places":
        return {
            "title": title,
            "headers": ["استان", "شهر", "تعداد سفارش", "فروش (تومان)"],
            "rows": [
                [row["province"], row["city"], row["orders"], row["revenue"]]
                for row in customers["places"]
            ],
            "numeric": {2, 3},
        }

    if key == "coupons":
        return {
            "title": title,
            "headers": ["کد کوپن", "تعداد استفاده", "مجموع تخفیف (تومان)"],
            "rows": [[row["code"], row["uses"], row["discount"]] for row in report["coupons"]],
            "numeric": {1, 2},
        }

    # chart
    return {
        "title": title,
        "headers": ["تاریخ (شمسی)", "سفارش‌ها", "فروش (تومان)"],
        "rows": [
            [_jalali(point["date"]), point["orders"], point["revenue"]]
            for point in report["chart"]["points"]
        ],
        "numeric": {1, 2},
    }


# ---------------------------------------------------------------------------
# writers
# ---------------------------------------------------------------------------

def xlsx_bytes(headers, rows, *, sheet_title="analytics", numeric_columns=()) -> tuple[bytes, int]:
    """Build an .xlsx workbook; returns (bytes, rows_written).

    The caller caps ``rows`` (see row_limit()); this function also stops
    writing at the cap as a second line of defence.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    limit = row_limit()
    headers = [str(header) for header in headers]
    numeric_columns = set(numeric_columns)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = (sheet_title or "analytics")[:31]
    sheet.sheet_view.rightToLeft = True

    sheet.append(headers)
    header_font = Font(bold=True, color="111827")
    header_fill = PatternFill("solid", fgColor="E8EAED")
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    sheet.freeze_panes = "A2"

    widths = [len(str(header)) for header in headers]
    written = 0
    for row in rows:
        if written >= limit:
            break
        sheet.append(list(row))
        written += 1
        for index, value in enumerate(row):
            if index in numeric_columns and isinstance(value, (int, float)):
                sheet.cell(row=written + 1, column=index + 1).number_format = "#,##0"
            if index < len(widths):
                widths[index] = max(widths[index], len(str(value if value is not None else "")))

    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 10), 45)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue(), written


def csv_bytes(headers, rows) -> bytes:
    """UTF-8 CSV with a BOM (Excel shows Persian text correctly)."""
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(headers)
    limit = row_limit()
    for index, row in enumerate(rows):
        if index >= limit:
            break
        writer.writerow(row)
    return ("\ufeff" + buffer.getvalue()).encode("utf-8")


def file_stem(prefix: str, range_dict: dict | None = None) -> str:
    if not range_dict:
        return prefix
    start = range_dict.get("start_date")
    end = range_dict.get("end_date")
    return f"{prefix}-{start}-{end}"


def xlsx_response(data: bytes, filename: str) -> HttpResponse:
    response = HttpResponse(data, content_type=XLSX_CONTENT_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def csv_response(data: bytes, filename: str) -> HttpResponse:
    response = HttpResponse(data, content_type=CSV_CONTENT_TYPE)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
