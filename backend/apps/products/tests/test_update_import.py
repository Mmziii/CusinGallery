"""
Spreadsheet price/stock update + export tests (Part R4 item 6).

Real behaviour: the engine functions run against real Product rows, the
admin views are exercised through real HTTP posts (multipart upload,
session-based preview -> confirm), and exports are parsed back with the
csv module / openpyxl to prove the format.
"""
import io

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse

from apps.categories.models import Category

from ..importing import (
    TEMPLATE_COLUMNS,
    export_products_csv,
    export_products_xlsx,
    update_products_from_rows,
)
from ..models import Product
from .helpers import make_product


def csv_upload(text, name="update.csv"):
    return SimpleUploadedFile(name, text.encode("utf-8-sig"), content_type="text/csv")


# Admin templates need {% static %}; keep tests manifest-independent
# (same reasoning as test_admin_products.py).
PLAIN_STATIC = {
    "STORAGES": {
        "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
        "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
    }
}


class UpdateEngineTests(TestCase):
    def setUp(self):
        self.category = Category.objects.create(name="Cookware", slug="cookware")
        self.product = make_product(
            category=self.category, name="ماگ سرامیکی", sku="MUG-1",
            price=100000, compare_at_price=200000, stock_quantity=10,
            short_description="توضیح اولیه",
        )
        self.header = ["sku", "price", "sale_price", "stock"]

    def run_rows(self, rows_text, dry_run=False):
        import csv as csv_mod

        reader = csv_mod.DictReader(io.StringIO(rows_text))
        rows = list(reader)
        return update_products_from_rows(rows, self.header, dry_run=dry_run)

    def test_updates_only_listed_fields_and_never_creates(self):
        before = Product.objects.count()
        result = self.run_rows("sku,price,sale_price,stock\nMUG-1,150000,,25\n")
        self.assertTrue(result.ok)
        self.assertTrue(result.applied)

        self.product.refresh_from_db()
        self.assertEqual(self.product.price, 150000)
        self.assertEqual(self.product.stock_quantity, 25)
        # untouched fields keep their values
        self.assertEqual(self.product.compare_at_price, 200000)
        self.assertEqual(self.product.name, "ماگ سرامیکی")
        self.assertEqual(self.product.short_description, "توضیح اولیه")
        self.assertEqual(Product.objects.count(), before)  # nothing created

    def test_unknown_sku_is_a_row_error_and_nothing_is_written(self):
        result = self.run_rows("sku,price,sale_price,stock\nNOPE-9,150000,,25\n")
        self.assertFalse(result.ok)
        self.assertIn("NOPE-9", result.errors[0])
        self.product.refresh_from_db()
        self.assertEqual(self.product.price, 100000)

    def test_dry_run_previews_old_to_new_without_writing(self):
        result = self.run_rows("sku,price,sale_price,stock\nMUG-1,150000,,25\n", dry_run=True)
        self.assertTrue(result.ok)
        self.assertFalse(result.applied)
        self.assertEqual(result.updated, 1)
        fields = {(c["field"], c["old"], c["new"]) for c in result.changes}
        self.assertIn(("price", 100000, 150000), fields)
        self.assertIn(("stock_quantity", 10, 25), fields)
        self.product.refresh_from_db()
        self.assertEqual(self.product.price, 100000)  # nothing written

    def test_unchanged_rows_are_counted_as_skipped(self):
        result = self.run_rows("sku,price,sale_price,stock\nMUG-1,100000,,10\n", dry_run=True)
        self.assertTrue(result.ok)
        self.assertEqual(result.updated, 0)
        self.assertEqual(result.skipped, 1)
        self.assertEqual(result.changes, [])

    def test_blank_or_non_numeric_price_is_a_row_error(self):
        result = self.run_rows("sku,price,sale_price,stock\nMUG-1,abc,,25\n")
        self.assertFalse(result.ok)
        self.assertIn("price", result.errors[0])

        result = self.run_rows("sku,price,sale_price,stock\n,150000,,25\n")
        self.assertFalse(result.ok)
        self.assertIn("sku", result.errors[0])

    def test_persian_digits_are_normalized(self):
        result = self.run_rows("sku,price,sale_price,stock\nMUG-1,۱۵۰۰۰۰,,۲۵\n")
        self.assertTrue(result.ok)
        self.product.refresh_from_db()
        self.assertEqual(self.product.price, 150000)
        self.assertEqual(self.product.stock_quantity, 25)

    def test_sale_price_below_price_is_rejected(self):
        result = self.run_rows("sku,price,sale_price,stock\nMUG-1,200000,150000,\n")
        self.assertFalse(result.ok)
        self.assertIn("sale_price", result.errors[0])

    def test_all_or_nothing_across_rows(self):
        # First row valid, second row bad SKU -> nothing must be written.
        result = self.run_rows(
            "sku,price,sale_price,stock\nMUG-1,150000,,25\nNOPE-9,1,,1\n"
        )
        self.assertFalse(result.ok)
        self.product.refresh_from_db()
        self.assertEqual(self.product.price, 100000)

    def test_unknown_column_is_rejected(self):
        import csv as csv_mod

        text = "sku,price,name\nMUG-1,150000,hack\n"
        rows = list(csv_mod.DictReader(io.StringIO(text)))
        result = update_products_from_rows(rows, ["sku", "price", "name"])
        self.assertFalse(result.ok)
        self.assertIn("name", result.errors[0])


class ExportTests(TestCase):
    def setUp(self):
        self.category = Category.objects.create(name="Cookware", slug="cookware")
        self.product = make_product(
            category=self.category, name="ماگ سرامیکی", sku="MUG-9",
            price=180000, compare_at_price=220000, stock_quantity=40,
        )

    def test_csv_export_uses_the_import_template_format(self):
        import csv as csv_mod

        text = export_products_csv()
        reader = csv_mod.DictReader(io.StringIO(text))
        self.assertEqual(reader.fieldnames, TEMPLATE_COLUMNS)
        rows = list(reader)
        self.assertEqual(len(rows), 1)
        row = rows[0]
        self.assertEqual(row["sku"], "MUG-9")
        self.assertEqual(row["name"], "ماگ سرامیکی")
        self.assertEqual(row["price"], "180000")
        self.assertEqual(row["compare_at_price"], "220000")
        self.assertEqual(row["stock_quantity"], "40")
        self.assertEqual(row["category_slug"], "cookware")

    def test_xlsx_export_is_readable_and_round_trips(self):
        from openpyxl import load_workbook

        data = export_products_xlsx()
        workbook = load_workbook(io.BytesIO(data), read_only=True)
        grid = list(workbook.worksheets[0].iter_rows(values_only=True))
        workbook.close()
        self.assertEqual(list(grid[0]), TEMPLATE_COLUMNS)
        row = grid[1]
        self.assertIn("MUG-9", row)
        self.assertIn("ماگ سرامیکی", row)

    def test_export_then_reimport_is_a_noop_update(self):
        """The owner's real loop: export, edit nothing, re-upload in
        update-only mode with the price/stock values -> no changes."""
        import csv as csv_mod

        text = export_products_csv()
        reader = csv_mod.DictReader(io.StringIO(text))
        full_rows = list(reader)
        minimal = [
            {"sku": r["sku"], "price": r["price"],
             "sale_price": r["compare_at_price"], "stock": r["stock_quantity"]}
            for r in full_rows
        ]
        result = update_products_from_rows(
            minimal, ["sku", "price", "sale_price", "stock"], dry_run=True
        )
        self.assertTrue(result.ok)
        self.assertEqual(result.updated, 0)
        self.assertEqual(result.skipped, 1)


@override_settings(**PLAIN_STATIC)
class AdminUpdateFlowTests(TestCase):
    """The real admin HTTP flow: upload -> preview (session) -> confirm."""

    def setUp(self):
        self.superuser = get_user_model().objects.create_superuser(
            username="owner-upd", password="x", phone="+989007778899"
        )
        self.client.force_login(self.superuser)
        self.category = Category.objects.create(name="Cookware", slug="cookware")
        self.product = make_product(
            category=self.category, name="کتری", sku="KET-7",
            price=100000, stock_quantity=5,
        )
        self.import_url = reverse("admin:products_product_import")

    def test_export_endpoint_returns_csv_attachment(self):
        response = self.client.get(reverse("admin:products_product_export"))
        self.assertEqual(response.status_code, 200)
        self.assertIn("attachment", response["Content-Disposition"])
        content = b"".join(response.streaming_content).decode("utf-8-sig")
        self.assertIn("KET-7", content)
        self.assertIn(",".join(TEMPLATE_COLUMNS), content)

    def test_export_endpoint_xlsx(self):
        response = self.client.get(
            reverse("admin:products_product_export"), {"format": "xlsx"}
        )
        self.assertEqual(response.status_code, 200)
        data = b"".join(response.streaming_content)
        self.assertTrue(data.startswith(b"PK"))  # zip (xlsx) magic bytes

    def test_update_only_preview_then_confirm_applies_once(self):
        upload = csv_upload("sku,price,sale_price,stock\nKET-7,140000,,9\n")
        response = self.client.post(
            self.import_url, {"file": upload, "update_only": "1"}
        )
        # Redirected to the preview page; nothing written yet.
        self.assertEqual(response.status_code, 302)
        self.assertIn("pending=1", response["Location"])
        self.product.refresh_from_db()
        self.assertEqual(self.product.price, 100000)

        preview = self.client.get(self.import_url, {"pending": "1"})
        self.assertEqual(preview.status_code, 200)
        html = preview.content.decode()
        self.assertIn("KET-7", html)
        self.assertIn("100000", html)   # old value
        self.assertIn("140000", html)   # new value

        response = self.client.post(self.import_url, {"confirm": "1"})
        self.assertEqual(response.status_code, 302)
        self.product.refresh_from_db()
        self.assertEqual(self.product.price, 140000)
        self.assertEqual(self.product.stock_quantity, 9)

        # Confirming again must be a no-op (the pending entry is gone).
        response = self.client.post(self.import_url, {"confirm": "1"}, follow=True)
        self.product.refresh_from_db()
        self.assertEqual(self.product.price, 140000)

    def test_update_only_with_errors_writes_nothing_and_shows_them(self):
        upload = csv_upload("sku,price,sale_price,stock\nGHOST-1,140000,,9\n")
        response = self.client.post(
            self.import_url, {"file": upload, "update_only": "1"}
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn("GHOST-1", response.content.decode())
        self.product.refresh_from_db()
        self.assertEqual(self.product.price, 100000)

    def test_full_import_mode_still_creates_new_products(self):
        text = (
            "name,slug,sku,category_slug,brand_name,price,compare_at_price,"
            "discount_percentage,stock_quantity,low_stock_threshold,"
            "short_description,description,is_active,is_featured\n"
            "قابلمه,ghablame,pot-1,cookware,,500000,,,7,,,,true,false\n"
        )
        upload = csv_upload(text, name="full.csv")
        response = self.client.post(self.import_url, {"file": upload})
        self.assertEqual(response.status_code, 302)
        self.assertTrue(Product.objects.filter(sku="pot-1").exists())
