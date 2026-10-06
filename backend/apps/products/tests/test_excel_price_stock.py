"""
Part S5 item 8: Excel price/stock workflow for the shop owner.

Real behaviour only: the engine runs against real Product rows, uploads
are real .xlsx/.csv bytes parsed by openpyxl/csv, and the admin flow is
driven through real multipart POSTs (sheet chooser -> column mapping ->
preview -> confirm) against the real templates.
"""
import csv as csv_mod
import io

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse

from apps.categories.models import Category

from ..importing import (
    UPDATE_TEMPLATE_COLUMNS,
    _sku_text,
    canonical_rows,
    errors_csv,
    normalize_header,
    read_rows,
    resolve_headers,
    sheet_names,
    update_products_from_rows,
)
from ..models import Brand, Product
from .helpers import make_product

PLAIN_STATIC = {
    "STORAGES": {
        "default": {"BACKEND": "django.core.files.storage.FileSystemStorage"},
        "staticfiles": {"BACKEND": "django.contrib.staticfiles.storage.StaticFilesStorage"},
    }
}


def xlsx_upload(sheets, name="products.xlsx"):
    """Build a real workbook upload: {sheet name: [row, row, ...]}."""
    from openpyxl import Workbook

    workbook = Workbook()
    first = True
    for title, grid in sheets.items():
        sheet = workbook.active if first else workbook.create_sheet()
        sheet.title = title
        for row in grid:
            sheet.append(row)
        first = False
    buffer = io.BytesIO()
    workbook.save(buffer)
    return SimpleUploadedFile(
        name,
        buffer.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def csv_upload(text, name="update.csv"):
    return SimpleUploadedFile(name, text.encode("utf-8-sig"), content_type="text/csv")


def rows_from_text(text):
    return list(csv_mod.DictReader(io.StringIO(text)))


class HeaderAliasTests(TestCase):
    """The owner's own column names, not ours."""

    def test_persian_aliases_resolve_to_canonical_columns(self):
        header = [
            "کد کالا", "نام کالا", "قیمت", "قیمت حراج", "موجودی",
            "دسته‌بندی", "برند", "وضعیت",
        ]
        mapping, unknown = resolve_headers(header, mode="update")
        self.assertEqual(unknown, [])
        self.assertEqual(mapping, {
            "کد کالا": "sku",
            "نام کالا": "name",
            "قیمت": "price",
            "قیمت حراج": "compare_at_price",
            "موجودی": "stock_quantity",
            "دسته‌بندی": "category_slug",
            "برند": "brand_name",
            "وضعیت": "is_active",
        })

    def test_english_short_forms_and_synonyms_work(self):
        mapping, unknown = resolve_headers(["SKU", "Name", "Price", "Stock", "Active"])
        self.assertEqual(unknown, [])
        self.assertEqual(
            [mapping[c] for c in ["SKU", "Name", "Price", "Stock", "Active"]],
            ["sku", "name", "price", "stock_quantity", "is_active"],
        )

    def test_case_spaces_zwnj_and_arabic_letters_are_ignored(self):
        # Arabic keyboard letters (ك/ي), a ZWNJ, extra spaces, mixed case.
        header = [" sku ", "قيمت‌فروش", "كدمحصول", "دسته بندی", "موجودی انبار"]
        mapping, unknown = resolve_headers(header, mode="update")
        self.assertEqual(unknown, [])
        self.assertEqual(mapping["كدمحصول"], "sku")
        self.assertEqual(mapping[" sku "], "sku")
        self.assertEqual(mapping["قيمت‌فروش"], "price")
        self.assertEqual(mapping["دسته بندی"], "category_slug")
        self.assertEqual(mapping["موجودی انبار"], "stock_quantity")
        self.assertEqual(normalize_header("قيمت‌فروش"), normalize_header("قیمت فروش"))

    def test_unknown_columns_are_reported_not_guessed(self):
        mapping, unknown = resolve_headers(["کد کالا", "قیمت", "colour", "توضیح من"], mode="update")
        self.assertEqual(unknown, ["colour", "توضیح من"])
        self.assertIsNone(mapping["colour"])

    def test_full_only_columns_are_ignored_in_the_update_flow(self):
        """An exported file re-uploaded for a price update must not require
        mapping: its extra columns are understood, just not writable."""
        mapping, unknown = resolve_headers(
            ["sku", "price", "stock_quantity", "slug", "short_description", "description",
             "discount_percentage", "low_stock_threshold", "is_featured"],
            mode="update",
        )
        self.assertEqual(unknown, [])
        self.assertIsNone(mapping["slug"])
        self.assertIsNone(mapping["short_description"])
        self.assertEqual(mapping["stock_quantity"], "stock_quantity")


class CanonicalRowTests(TestCase):
    def test_rows_are_rekeyed_and_leading_zeros_survive(self):
        rows = rows_from_text("کد کالا,قیمت,موجودی\n00123,\"۱٬۲۳۴٬۰۰۰\", 7 \n")
        mapping, unknown = resolve_headers(["کد کالا", "قیمت", "موجودی"], mode="update")
        self.assertEqual(unknown, [])
        converted, header = canonical_rows(rows, mapping, mode="update")
        self.assertEqual(set(header), {"sku", "price", "stock"})
        self.assertEqual(converted[0]["sku"], "00123")
        self.assertEqual(converted[0]["price"], "۱٬۲۳۴٬۰۰۰")


class NumberToleranceTests(TestCase):
    def setUp(self):
        self.category = Category.objects.create(name="Cookware", slug="cookware")
        self.product = make_product(
            category=self.category, name="کتری", sku="KET-7", price=100000, stock_quantity=5,
        )

    def run_update(self, text, header=("sku", "price", "sale_price", "stock"), **kwargs):
        rows = rows_from_text(text)
        return update_products_from_rows(rows, list(header), **kwargs)

    def test_thousand_separators_persian_digits_and_spaces(self):
        result = self.run_update(
            "sku,price,stock\nKET-7,\"۱٬۲۳۴٬۰۰۰\",\" ۹ \"\n"
        )
        self.assertTrue(result.ok, result.errors)
        self.product.refresh_from_db()
        self.assertEqual(self.product.price, 1234000)   # ۱۲۳۴۰۰۰ with a Persian thousands mark
        self.assertEqual(self.product.stock_quantity, 9)

    def test_ascii_thousand_separators_and_arabic_digits(self):
        result = self.run_update("sku,price,stock\nKET-7,\"1,234,000\",٣\n")
        self.assertTrue(result.ok, result.errors)
        self.product.refresh_from_db()
        self.assertEqual(self.product.price, 1234000)
        self.assertEqual(self.product.stock_quantity, 3)

    def test_excel_float_cells_are_accepted(self):
        result = self.run_update("sku,price,stock\nKET-7,190000.0,12.0\n")
        self.assertTrue(result.ok, result.errors)
        self.product.refresh_from_db()
        self.assertEqual(self.product.price, 190000)
        self.assertEqual(self.product.stock_quantity, 12)

    def test_a_fractional_price_is_still_an_error(self):
        result = self.run_update("sku,price\nKET-7,190000.5\n")
        self.assertFalse(result.ok)
        self.assertIn("عدد صحیح", result.errors[0])

    def test_empty_rows_are_ignored(self):
        upload = xlsx_upload({"Sheet1": [
            list(UPDATE_TEMPLATE_COLUMNS),
            ["KET-7", 150000, None, 8],
            [None, None, None, None],
            ["", "", "", ""],
        ]})
        rows, header = read_rows(upload)
        self.assertEqual(len(rows), 1)
        self.assertEqual(header[0], "sku")


class SkuTextTests(TestCase):
    def setUp(self):
        self.category = Category.objects.create(name="Cookware", slug="cookware")
        self.zero = make_product(
            category=self.category, name="کالای صفر ابتدایی", sku="00123", price=50000,
        )
        self.integer = make_product(
            category=self.category, name="کالای عددی", sku="1234", price=60000,
        )

    def test_sku_with_leading_zeros_stays_text(self):
        self.assertEqual(_sku_text("00123"), "00123")
        result = update_products_from_rows(
            rows_from_text("sku,price\n00123,77000\n"), ["sku", "price"]
        )
        self.assertTrue(result.ok, result.errors)
        self.zero.refresh_from_db()
        self.assertEqual(self.zero.price, 77000)
        self.assertFalse(Product.objects.filter(sku="123").exists())

    def test_numeric_excel_cell_1234_point_0_becomes_1234(self):
        upload = xlsx_upload({"Sheet1": [["sku", "price"], [1234.0, 88000]]})
        rows, header = read_rows(upload)
        self.assertEqual(_sku_text(rows[0]["sku"]), "1234")
        result = update_products_from_rows(rows, header)
        self.assertTrue(result.ok, result.errors)
        self.integer.refresh_from_db()
        self.assertEqual(self.integer.price, 88000)


class StockModeTests(TestCase):
    def setUp(self):
        self.category = Category.objects.create(name="Cookware", slug="cookware")
        self.product = make_product(
            category=self.category, name="کتری", sku="KET-7", price=100000, stock_quantity=10,
        )

    def test_replace_mode_sets_the_stock(self):
        result = update_products_from_rows(
            rows_from_text("sku,stock\nKET-7,4\n"), ["sku", "stock"]
        )
        self.assertTrue(result.ok, result.errors)
        self.product.refresh_from_db()
        self.assertEqual(self.product.stock_quantity, 4)

    def test_add_mode_adds_the_delta(self):
        result = update_products_from_rows(
            rows_from_text("sku,stock\nKET-7,15\n"), ["sku", "stock"], stock_mode="add"
        )
        self.assertTrue(result.ok, result.errors)
        self.product.refresh_from_db()
        self.assertEqual(self.product.stock_quantity, 25)

    def test_add_mode_accepts_a_negative_delta(self):
        result = update_products_from_rows(
            rows_from_text("sku,stock\nKET-7,-4\n"), ["sku", "stock"], stock_mode="add"
        )
        self.assertTrue(result.ok, result.errors)
        self.product.refresh_from_db()
        self.assertEqual(self.product.stock_quantity, 6)

    def test_add_mode_never_goes_below_zero(self):
        result = update_products_from_rows(
            rows_from_text("sku,stock\nKET-7,-50\n"), ["sku", "stock"],
            stock_mode="add", dry_run=True,
        )
        self.assertFalse(result.ok)
        self.assertEqual(self.product.stock_quantity, 10)

    def test_an_unknown_mode_is_rejected(self):
        result = update_products_from_rows(
            rows_from_text("sku,stock\nKET-7,1\n"), ["sku", "stock"], stock_mode="multiply"
        )
        self.assertFalse(result.ok)
        self.assertIn("حالت موجودی", result.errors[0])


class SheetChooserTests(TestCase):
    def setUp(self):
        self.category = Category.objects.create(name="Cookware", slug="cookware")
        self.product = make_product(
            category=self.category, name="کتری", sku="KET-7", price=100000, stock_quantity=10,
        )

    def test_sheet_names_are_listed(self):
        upload = xlsx_upload({"اول": [["sku", "price"]], "دوم": [["sku", "price"]]})
        self.assertEqual(sheet_names(upload), ["اول", "دوم"])

    def test_an_unknown_sheet_name_is_a_clear_error(self):
        upload = xlsx_upload({"اول": [["sku", "price"], ["KET-7", 1]]})
        with self.assertRaises(ValueError) as ctx:
            read_rows(upload, sheet="سوم")
        self.assertIn("اول", str(ctx.exception))

    def test_reading_the_chosen_sheet(self):
        upload = xlsx_upload({
            "اول": [["sku", "price"], ["KET-7", 111111]],
            "دوم": [["sku", "price"], ["KET-7", 222222]],
        })
        rows, _header = read_rows(upload, sheet="دوم")
        self.assertEqual(rows[0]["price"], 222222)

    def test_the_engine_does_not_care_which_sheet_was_chosen(self):
        upload = xlsx_upload({
            "اول": [["sku", "price"], ["KET-7", 111111]],
            "دوم": [["sku", "price"], ["KET-7", 999999]],
        })
        rows, header = read_rows(upload, sheet="دوم")
        result = update_products_from_rows(rows, header)
        self.assertTrue(result.ok, result.errors)
        self.product.refresh_from_db()
        self.assertEqual(self.product.price, 999999)


class OldXlsTests(TestCase):
    def test_xls_gets_a_clear_persian_message(self):
        upload = SimpleUploadedFile("old.xls", b"\xd0\xcf\x11\xe0", content_type="application/vnd.ms-excel")
        with self.assertRaises(ValueError) as ctx:
            read_rows(upload)
        message = str(ctx.exception)
        self.assertIn(".xls", message)
        self.assertIn(".xlsx", message)


class ErrorReportTests(TestCase):
    def test_error_csv_has_a_row_number_column(self):
        text = errors_csv([
            "ردیف 3: ستون «price» -- «abc» یک عدد صحیح نیست.",
            "ستون(های) ناشناخته: colour.",
        ])
        rows = list(csv_mod.reader(io.StringIO(text)))
        self.assertEqual(rows[0], ["ردیف", "خطا"])
        self.assertEqual(rows[1][0], "3")
        self.assertIn("price", rows[1][1])
        self.assertEqual(rows[2][0], "")
        self.assertIn("colour", rows[2][1])


@override_settings(**PLAIN_STATIC)
class AdminExcelWorkflowTests(TestCase):
    """The three buttons and the whole upload -> steps -> apply flow."""

    def setUp(self):
        self.superuser = get_user_model().objects.create_superuser(
            username="owner-excel", password="x", phone="+989001112233"
        )
        self.client.force_login(self.superuser)
        self.category = Category.objects.create(name="Cookware", slug="cookware")
        self.brand = Brand.objects.create(name="کازین", slug="kazin")
        self.product = make_product(
            category=self.category, name="کتری", sku="KET-7",
            price=100000, stock_quantity=10,
        )
        self.import_url = reverse("admin:products_product_import")

    def test_changelist_shows_the_three_actions_each_with_a_template(self):
        response = self.client.get(reverse("admin:products_product_changelist"))
        html = response.content.decode()
        self.assertIn("ورود محصولات از اکسل", html)
        self.assertIn("به‌روزرسانی قیمت و موجودی", html)
        self.assertIn("خروجی اکسل", html)
        # Each action links to its own downloadable template.
        self.assertIn(reverse("admin:products_product_import_template") + "?format=xlsx", html)
        self.assertIn(reverse("admin:products_product_update_template") + "?format=xlsx", html)
        self.assertIn(reverse("admin:products_product_export") + "?format=xlsx", html)

    def test_templates_download_as_xlsx_and_parse_back(self):
        from openpyxl import load_workbook

        full = self.client.get(
            reverse("admin:products_product_import_template"), {"format": "xlsx"}
        )
        data = b"".join(full.streaming_content)
        self.assertTrue(data.startswith(b"PK"))
        sheet = load_workbook(io.BytesIO(data)).active
        header = [cell.value for cell in sheet[1]]
        self.assertEqual(header[:5], ["name", "slug", "sku", "category_slug", "brand_name"])
        self.assertEqual(len(header), 14)

    def test_update_via_persian_headers_end_to_end(self):
        upload = csv_upload("کد کالا,قیمت,موجودی\nKET-7,\"۱۹۰٬۰۰۰\",4\n")
        response = self.client.post(self.import_url, {"file": upload, "mode": "update"})
        self.assertEqual(response.status_code, 302)
        self.assertIn("pending=1", response["Location"])
        self.product.refresh_from_db()
        self.assertEqual(self.product.price, 100000)   # preview only

        preview = self.client.get(self.import_url, {"pending": "1"})
        html = preview.content.decode()
        self.assertIn("KET-7", html)
        self.assertIn("190000", html)

        self.client.post(self.import_url, {"confirm": "1"})
        self.product.refresh_from_db()
        self.assertEqual(self.product.price, 190000)
        self.assertEqual(self.product.stock_quantity, 4)

    def test_add_stock_mode_through_the_form(self):
        upload = csv_upload("sku,stock\nKET-7,5\n")
        self.client.post(
            self.import_url, {"file": upload, "mode": "update", "stock_mode": "add"}
        )
        self.client.post(self.import_url, {"confirm": "1"})
        self.product.refresh_from_db()
        self.assertEqual(self.product.stock_quantity, 15)

    def test_unknown_column_asks_for_a_mapping_with_a_preview(self):
        upload = csv_upload("کد کالا,قیمت,تعداد بسته\nKET-7,150000,3\n")
        response = self.client.post(self.import_url, {"file": upload, "mode": "update"})
        self.assertEqual(response.status_code, 200)
        html = response.content.decode()
        # The dropdown + a preview of the real values, plus the token that
        # lets the next step reuse the uploaded file.
        self.assertIn("تعداد بسته", html)
        self.assertIn('<select name="map">', html)
        self.assertIn("150000", html)
        self.assertIn('name="token"', html)
        self.assertIn("نادیده بگیر", html)
        self.product.refresh_from_db()
        self.assertEqual(self.product.price, 100000)

    def test_mapping_an_unknown_column_to_a_field_applies_it(self):
        upload = csv_upload("کد کالا,قیمت,تعداد بسته\nKET-7,150000,3\n")
        first = self.client.post(self.import_url, {"file": upload, "mode": "update"})
        token = self._token_from(first.content.decode())
        response = self.client.post(self.import_url, {
            "mode": "update",
            "token": token,
            "step": "mapping",
            "unknown_columns": ["تعداد بسته"],
            "map": ["stock_quantity"],
        })
        self.assertEqual(response.status_code, 302)
        self.client.post(self.import_url, {"confirm": "1"})
        self.product.refresh_from_db()
        self.assertEqual(self.product.price, 150000)
        self.assertEqual(self.product.stock_quantity, 3)

    def test_ignoring_an_unknown_column_leaves_the_product_alone(self):
        upload = csv_upload("کد کالا,تعداد بسته\nKET-7,3\n")
        first = self.client.post(self.import_url, {"file": upload, "mode": "update"})
        token = self._token_from(first.content.decode())
        # "Ignore" on the only data column leaves nothing to update.
        response = self.client.post(self.import_url, {
            "mode": "update", "token": token, "step": "mapping",
            "unknown_columns": ["تعداد بسته"], "map": [""],
        }, follow=True)
        self.assertEqual(response.status_code, 200)
        self.product.refresh_from_db()
        self.assertEqual(self.product.stock_quantity, 10)

    def test_workbook_with_several_sheets_asks_which_one(self):
        upload = xlsx_upload({
            "راهنما": [["این شیت داده نیست"]],
            "محصولات": [["sku", "price"], ["KET-7", 175000]],
        })
        response = self.client.post(self.import_url, {"file": upload, "mode": "update"})
        html = response.content.decode()
        self.assertIn("کدام شیت", html)
        self.assertIn("محصولات", html)
        self.assertIn("راهنما", html)
        self.product.refresh_from_db()
        self.assertEqual(self.product.price, 100000)

        token = self._token_from(html)
        response = self.client.post(self.import_url, {
            "mode": "update", "token": token, "sheet": "محصولات",
        })
        self.assertEqual(response.status_code, 302)
        self.client.post(self.import_url, {"confirm": "1"})
        self.product.refresh_from_db()
        self.assertEqual(self.product.price, 175000)

    def test_old_xls_upload_shows_the_persian_message(self):
        upload = SimpleUploadedFile("old.xls", b"\xd0\xcf\x11\xe0", content_type="application/vnd.ms-excel")
        response = self.client.post(self.import_url, {"file": upload, "mode": "update"})
        html = response.content.decode()
        self.assertIn("xls", html)
        self.assertIn(".xlsx", html)

    def test_errors_can_be_downloaded_as_a_per_row_csv(self):
        upload = csv_upload("sku,price\nKET-7,not-a-number\nGHOST-1,1000\n")
        response = self.client.post(self.import_url, {"file": upload, "mode": "update"})
        self.assertEqual(response.status_code, 200)
        self.assertIn("دانلود فایل خطاها", response.content.decode())

        report = self.client.get(self.import_url, {"errors": "csv"})
        self.assertEqual(report.status_code, 200)
        text = report.content.decode("utf-8-sig")
        rows = list(csv_mod.reader(io.StringIO(text)))
        self.assertEqual(rows[0], ["ردیف", "خطا"])
        self.assertEqual(rows[1][0], "2")
        self.assertEqual(rows[2][0], "3")
        self.product.refresh_from_db()
        self.assertEqual(self.product.price, 100000)

    def test_duplicate_sku_in_the_file_is_a_row_error(self):
        upload = csv_upload("sku,price\nKET-7,150000\nKET-7,160000\n")
        response = self.client.post(self.import_url, {"file": upload, "mode": "update"})
        self.assertEqual(response.status_code, 200)
        self.assertIn("بیش از یک بار", response.content.decode())
        self.product.refresh_from_db()
        self.assertEqual(self.product.price, 100000)

    def test_full_import_accepts_persian_headers_and_creates_products(self):
        upload = csv_upload(
            "نام کالا,شناسه,کد کالا,دسته‌بندی,برند,قیمت,موجودی\n"
            "قابلمه گرانیتی,ghablame,GA-1,cookware,کازین,۵۰۰۰۰۰,۷\n",
            name="full.csv",
        )
        response = self.client.post(self.import_url, {"file": upload, "mode": "full"})
        self.assertEqual(response.status_code, 302)
        product = Product.objects.get(sku="GA-1")
        self.assertEqual(product.price, 500000)
        self.assertEqual(product.stock_quantity, 7)
        self.assertEqual(product.category, self.category)
        self.assertEqual(product.brand, self.brand)

    def test_full_import_dry_run_reports_counts_and_writes_nothing(self):
        upload = csv_upload(
            "name,slug,sku,category_slug,price\nقابلمه,ghablame,GA-2,cookware,500000\n",
            name="dry.csv",
        )
        response = self.client.post(
            self.import_url, {"file": upload, "mode": "full", "dry_run": "1"}
        )
        self.assertEqual(response.status_code, 200)
        html = response.content.decode()
        self.assertIn("هنوز چیزی ذخیره نشده", html)
        self.assertFalse(Product.objects.filter(sku="GA-2").exists())

    def test_optional_columns_update_name_category_brand_and_status(self):
        upload = csv_upload(
            "کد کالا,نام کالا,دسته‌بندی,برند,وضعیت\n"
            "KET-7,کتری جدید,cookware,کازین,خیر\n"
        )
        self.client.post(self.import_url, {"file": upload, "mode": "update"})
        self.client.post(self.import_url, {"confirm": "1"})
        self.product.refresh_from_db()
        self.assertEqual(self.product.name, "کتری جدید")
        self.assertEqual(self.product.category, self.category)
        self.assertEqual(self.product.brand, self.brand)
        self.assertFalse(self.product.is_active)

    def test_full_import_still_reports_missing_required_columns(self):
        upload = csv_upload("sku,price\nGA-3,1000\n", name="bad.csv")
        response = self.client.post(self.import_url, {"file": upload, "mode": "full"})
        self.assertEqual(response.status_code, 200)
        html = response.content.decode()
        self.assertIn("name", html)
        self.assertFalse(Product.objects.filter(sku="GA-3").exists())

    @staticmethod
    def _token_from(html):
        match = None
        for chunk in html.split('name="token" value="')[1:]:
            match = chunk.split('"')[0]
            if match:
                break
        assert match, "no stash token in the page"
        return match
